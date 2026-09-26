import io
from typing import Any, Iterable, Iterator

from openpyxl import Workbook


class ExcelReportStrategy:
    def generate(
        self,
        headers: list[str],
        data: Iterable[dict[str, Any]],
    ) -> Iterator[bytes]:
        workbook = Workbook()
        worksheet = workbook.active

        if worksheet is not None:
            worksheet.append(headers)

            for row in data:
                worksheet.append([row.get(header, "") for header in headers])

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        yield output.getvalue()


class MultiSheetExcelReportStrategy:
    def generate_multi_sheet(
        self,
        sheets_data: dict[str, dict[str, Any]],
    ) -> Iterator[bytes]:
        workbook = Workbook()
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)

        for sheet_name, content in sheets_data.items():
            worksheet = workbook.create_sheet(title=sheet_name)
            headers = content.get("headers", [])
            data = content.get("data", [])

            worksheet.append(headers)

            for row in data:
                worksheet.append([row.get(header, "") for header in headers])

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        yield output.getvalue()

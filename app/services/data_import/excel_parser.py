from openpyxl import load_workbook

from app.services.data_import.types import RawImportRow


class ExcelFileParser:
    def parse(self, file_path: str) -> list[RawImportRow]:
        workbook = load_workbook(
            filename=file_path,
            read_only=True,
            data_only=True,
        )

        try:
            worksheet = workbook.active

            rows = worksheet.iter_rows(values_only=True)

            headers = next(rows, None)

            if not headers:
                return []

            normalized_headers = [
                str(header).strip() if header is not None else "" for header in headers
            ]

            result = []

            for row in rows:
                result.append(
                    {
                        header: str(value) if value is not None else ""
                        for header, value in zip(normalized_headers, row)
                        if header
                    }
                )

            return result

        finally:
            workbook.close()

import io
from pathlib import Path
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.reporting.csv_strategy import CsvReportStrategy
from app.services.reporting.excel_strategy import (
    ExcelReportStrategy,
    MultiSheetExcelReportStrategy,
)
from app.services.reporting.generator import ReportGenerator
from app.worker.celery_app import celery_app


def test_csv_strategy_generation():
    strategy = CsvReportStrategy()
    headers = ["id", "name"]
    data = [{"id": 1, "name": "Amir"}, {"id": 2, "name": "Reza"}]

    generator = ReportGenerator(strategy)
    result = list(generator.generate(headers, data))

    assert len(result) == 3
    assert "id,name" in result[0].strip()
    assert "1,Amir" in result[1].strip()
    assert "2,Reza" in result[2].strip()


def test_excel_strategy_generation():
    strategy = ExcelReportStrategy()
    headers = ["id", "name"]
    data = [{"id": 1, "name": "Amir"}, {"id": 2, "name": "Reza"}]

    generator = ReportGenerator(strategy)
    result_bytes = b"".join(list(generator.generate(headers, data)))

    workbook = load_workbook(filename=io.BytesIO(result_bytes))
    worksheet = workbook.active

    rows = list(worksheet.iter_rows(values_only=True))
    assert rows[0] == ("id", "name")
    assert rows[1] == (1, "Amir")
    assert rows[2] == (2, "Reza")


def test_multi_sheet_excel_strategy_generation():
    strategy = MultiSheetExcelReportStrategy()
    sheets_data = {
        "Sheet1": {"headers": ["id", "name"], "data": [{"id": 1, "name": "Amir"}]},
        "Sheet2": {"headers": ["value"], "data": [{"value": 100}]},
    }

    result_bytes = b"".join(list(strategy.generate_multi_sheet(sheets_data)))
    workbook = load_workbook(filename=io.BytesIO(result_bytes))

    assert "Sheet1" in workbook.sheetnames
    assert "Sheet2" in workbook.sheetnames

    ws1 = workbook["Sheet1"]
    rows1 = list(ws1.iter_rows(values_only=True))
    assert rows1[0] == ("id", "name")
    assert rows1[1] == (1, "Amir")

    ws2 = workbook["Sheet2"]
    rows2 = list(ws2.iter_rows(values_only=True))
    assert rows2[0] == ("value",)
    assert rows2[1] == (100,)


def test_export_analytics_endpoints_csv(client: TestClient):
    response = client.get(
        "/analytics/products/performance/export?limit=10&export_format=csv"
    )
    assert response.status_code == 200
    assert "job_id" in response.json()

    response_categories = client.get(
        "/analytics/categories/performance/export?limit=10&export_format=csv"
    )
    assert response_categories.status_code == 200
    assert "job_id" in response_categories.json()

    response_trends = client.get(
        "/analytics/trends/export?period=daily&export_format=csv"
    )
    assert response_trends.status_code == 200
    assert "job_id" in response_trends.json()


def test_export_analytics_endpoints_excel(client: TestClient):
    response_trends = client.get(
        "/analytics/trends/export?period=daily&export_format=excel"
    )
    assert response_trends.status_code == 200
    assert "job_id" in response_trends.json()


def test_export_dashboard_excel(client: TestClient):
    response = client.get("/analytics/dashboard/export")
    assert response.status_code == 200
    assert "job_id" in response.json()


def test_async_report_status_and_download(client: TestClient, tmp_path: Path):
    dummy_file = tmp_path / "test_report.csv"
    dummy_file.write_text("id,name\n1,ProductA\n", encoding="utf-8")

    mock_pending = MagicMock()
    mock_pending.state = "PENDING"

    with patch.object(celery_app, "AsyncResult", return_value=mock_pending):
        status_res = client.get("/reports/status/job-pending")
        assert status_res.status_code == 200
        assert status_res.json() == {"status": "pending"}

        dl_res = client.get("/reports/download/job-pending")
        assert dl_res.status_code == 400

    mock_success = MagicMock()
    mock_success.state = "SUCCESS"
    mock_success.result = {"status": "completed", "file_path": str(dummy_file)}

    with patch.object(celery_app, "AsyncResult", return_value=mock_success):
        status_res = client.get("/reports/status/job-success")
        assert status_res.status_code == 200
        assert status_res.json() == {
            "status": "completed",
            "download_url": "/reports/download/job-success",
        }

        dl_res = client.get("/reports/download/job-success")
        assert dl_res.status_code == 200
        assert "id,name" in dl_res.text

    mock_missing_file = MagicMock()
    mock_missing_file.state = "SUCCESS"
    mock_missing_file.result = {
        "status": "completed",
        "file_path": str(tmp_path / "nonexistent.csv"),
    }

    with patch.object(celery_app, "AsyncResult", return_value=mock_missing_file):
        dl_res = client.get("/reports/download/job-missing")
        assert dl_res.status_code == 404


def test_concurrent_report_downloads(client: TestClient, tmp_path: Path):
    import concurrent.futures

    def mock_async_result(job_id):
        dummy_file = tmp_path / f"{job_id}.csv"
        dummy_file.write_text("id,val\n1,100\n", encoding="utf-8")

        mock_res = MagicMock()
        mock_res.state = "SUCCESS"
        mock_res.result = {"status": "completed", "file_path": str(dummy_file)}
        return mock_res

    def trigger_download(job_id):
        res = client.get(f"/reports/download/{job_id}")
        return res, tmp_path / f"{job_id}.csv"

    with patch(
        "app.routers.reports.celery_app.AsyncResult", side_effect=mock_async_result
    ):
        with concurrent.futures.ThreadPoolExecutor(max_workers=5) as executor:
            futures = [
                executor.submit(trigger_download, f"job-stress-{i}") for i in range(10)
            ]

            for future in concurrent.futures.as_completed(futures):
                res, file_path = future.result()
                assert res.status_code == 200
                assert "id,val" in res.text
                assert not file_path.exists(), "Temporary file was not cleaned up!"
